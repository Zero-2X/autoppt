"""Exercise the bundled builders through their real subprocess entrypoints."""

import json
import shutil
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from autoppt_workflow.ppt import pptx_builder as workflow_builder
from autopptskills.scripts.ppt import pptx_builder as skill_builder


@pytest.mark.parametrize("builder", [workflow_builder, skill_builder])
def test_builder_resolves_local_script_without_git_or_parent_project(tmp_path, monkeypatch, builder):
    # Emulate a downloaded source archive: no .git and no old sibling project.
    archive = tmp_path / "source archive"
    root = Path(__file__).resolve().parents[1]
    module_path = Path(builder.__file__).resolve().relative_to(root)
    monkeypatch.setattr(builder, "__file__", str(archive / module_path))
    script = archive / "autopptskills" / "scripts" / "build_image_deck.py"
    script.parent.mkdir(parents=True)
    shutil.copy2(root / "autopptskills" / "scripts" / "build_image_deck.py", script)

    workspace = tmp_path / "mock workspace"
    workspace.mkdir()
    # Structural fixture only; no provenance or release claim is made.
    Image.new("RGB", (160, 90), "white").save(workspace / "fixture.png")
    builder.build_deck_spec(workspace, {"slides": [{
        "slide_id": "S01", "final_path": "fixture.png", "headline": "Fixture",
    }]})
    output, report = workspace / "fixture.pptx", workspace / "build-report.json"
    builder.build_pptx(workspace, output, report)

    deck = Presentation(output)
    assert len(deck.slides) == 1
    assert len(deck.slides[0].shapes) == 1
    assert not any(shape.has_text_frame for shape in deck.slides[0].shapes)
    assert json.loads(report.read_text(encoding="utf-8"))["slide_count"] == 1
