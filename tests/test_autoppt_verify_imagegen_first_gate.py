from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "autopptskills" / "scripts" / "verify_imagegen_first_gate.py"


def load_module():
    spec = importlib.util.spec_from_file_location("autoppt_verify_imagegen_first_gate", SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class VerifyImagegenFirstGateTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.module = load_module()

    def test_accepts_explicit_chinese_complete_slide_contract(self) -> None:
        slide = {
            "prompt_zh": "生成一张完整、最终可用的中文PowerPoint幻灯片图片。画布：16:9。",
            "acceptance_criteria": ["逐字正确"],
        }

        self.assertTrue(self.module.has_full_slide_prompt_contract(slide))

    def test_accepts_explicit_english_complete_slide_contract(self) -> None:
        slide = {
            "prompt": "16:9 horizontal, 1920x1080, complete final PowerPoint slide image.",
            "expected_output": "complete final PPT page image",
        }

        self.assertTrue(self.module.has_full_slide_prompt_contract(slide))

    def test_rejects_generic_chinese_image_request(self) -> None:
        slide = {
            "prompt_zh": "生成一张中文图片。",
            "acceptance_criteria": ["逐字正确"],
        }

        self.assertFalse(self.module.has_full_slide_prompt_contract(slide))

    def official_reference(self) -> dict:
        return {
            "reference_mode": "official-source-inspired-generic-archetype",
            "official_source_inspired": True,
            "official_template": False,
            "sources": [{"url": "https://example.edu/template", "asset_reuse_allowed": False}],
            "allowed_extraction": ["generic hierarchy"],
            "forbidden_identity_elements": ["university_emblem_or_seal"],
            "identity_guard": "not an official template; no university emblem",
            "identity_text_policy": "source_verified_exact_text_only_not_identity_asset",
            "verified_identity_text": ["Example University"],
            "allowed_identity_text": ["Example University"],
        }

    def official_contract(self) -> tuple[dict, dict, list[dict], dict[str, dict]]:
        reference = self.official_reference()
        deck = {
            "style_profile": "graduate_defense_navy",
            "style_reference": reference,
        }
        manifest = {
            "style_profile": "graduate_defense_navy",
            "style_reference": reference,
            "prompt_policy": dict(self.module.OFFICIAL_SOURCE_REQUIRED_POLICY),
        }
        spec_slides = [{"slide_id": "S01"}]
        prompt_slides = {
            "S01": {
                "slide_id": "S01",
                "style_profile": "graduate_defense_navy",
                "style_reference": reference,
                "exact_text": ["Example University"],
            }
        }
        return deck, manifest, spec_slides, prompt_slides

    def test_official_source_contract_requires_policy_and_matching_slide_reference(self) -> None:
        deck, manifest, spec_slides, prompt_slides = self.official_contract()
        report, errors = self.module.validate_official_source_contract(
            deck, manifest, spec_slides, prompt_slides
        )
        self.assertEqual(errors, [])
        self.assertEqual(report["status"], "pass")
        self.assertTrue(report["required"])

        manifest["prompt_policy"].pop("forbid_official_template_claim")
        prompt_slides["S01"]["style_reference"] = {
            **prompt_slides["S01"]["style_reference"],
            "official_template": True,
        }
        report, errors = self.module.validate_official_source_contract(
            deck, manifest, spec_slides, prompt_slides
        )
        self.assertEqual(report["status"], "fail")
        self.assertTrue(any("forbid_official_template_claim" in error for error in errors))
        self.assertTrue(any("S01: style_reference differs" in error for error in errors))

    def test_official_source_contract_malformed_sources_fail_without_exception(self) -> None:
        deck, manifest, spec_slides, prompt_slides = self.official_contract()
        malformed_reference = {
            **manifest["style_reference"],
            "sources": [None, "not-an-object"],
        }
        manifest["style_reference"] = malformed_reference
        deck["style_reference"] = malformed_reference
        prompt_slides["S01"]["style_reference"] = malformed_reference

        report, errors = self.module.validate_official_source_contract(
            deck, manifest, spec_slides, prompt_slides
        )

        self.assertEqual(report["status"], "fail")
        self.assertTrue(
            any("asset_reuse_allowed=false" in error for error in errors),
            errors,
        )

    def test_full_official_source_gate_passes_an_explicit_fixture(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            slides_dir = root / "assets" / "slides"
            slides_dir.mkdir(parents=True)
            image_path = slides_dir / "S01.png"
            Image.new("RGB", (160, 90), "white").save(image_path)
            image_sha256 = hashlib.sha256(image_path.read_bytes()).hexdigest()

            reference = self.official_reference()
            deck_spec = {
                "deck": {
                    "style_profile": "graduate_defense_navy",
                    "style_reference": reference,
                    "output_mode": self.module.REQUIRED_DECK_MODE,
                    "generation_mode": self.module.REQUIRED_GENERATION_MODE,
                },
                "slides": [
                    {
                        "slide_id": "S01",
                        "layout": "full_slide_image",
                        "slide_image": "assets/slides/S01.png",
                    }
                ],
            }
            prompt_policy = {
                **self.module.REQUIRED_POLICY,
                **self.module.OFFICIAL_SOURCE_REQUIRED_POLICY,
            }
            prompt_manifest = {
                "style_profile": "graduate_defense_navy",
                "style_reference": reference,
                "prompt_policy": prompt_policy,
                "slides": [
                    {
                        "slide_id": "S01",
                        "style_profile": "graduate_defense_navy",
                        "style_reference": reference,
                        "exact_text": ["Example University"],
                        "prompt_zh": "生成一张完整、最终可用的中文PowerPoint幻灯片图片。",
                    }
                ],
            }
            (root / "deck-spec.json").write_text(
                json.dumps(deck_spec), encoding="utf-8"
            )
            (root / "image-prompts.json").write_text(
                json.dumps(prompt_manifest), encoding="utf-8"
            )
            provenance_path = root / "fixture-provenance.json"
            provenance_path.write_text(
                json.dumps(
                    {
                        "fixture_only": True,
                        "generation_mode": self.module.REQUIRED_GENERATION_MODE,
                        "slides": [
                            {
                                "sha256": image_sha256,
                                "backend": "builtin",
                                "status": "generated",
                                "provenance_kind": "imagegen-fixture-only",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )

            pptx_path = root / "fixture-image-only.pptx"
            prs = Presentation()
            prs.slide_width = 13_333_333
            prs.slide_height = 7_500_000
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )
            prs.save(pptx_path)

            report = self.module.run_gate(
                argparse.Namespace(
                    run_dir=str(root),
                    deck_spec="",
                    prompt_manifest="",
                    pptx=str(pptx_path),
                    provenance=[str(provenance_path)],
                    require_strong_ig_id=False,
                    report="",
                )
            )

            self.assertEqual(report["verdict"], "pass", report["errors"])
            self.assertEqual(report["slides_verified"], 1)
            self.assertEqual(report["assembly"]["full_slide_pictures"], 1)
            self.assertEqual(report["official_source_contract"]["status"], "pass")

    def test_non_official_source_contract_remains_compatible(self) -> None:
        report, errors = self.module.validate_official_source_contract(
            {"style_profile": "academic_light"},
            {"style_profile": "academic_light", "prompt_policy": {}},
            [{"slide_id": "S01"}],
            {"S01": {"slide_id": "S01", "style_profile": "academic_light"}},
        )
        self.assertEqual(errors, [])
        self.assertFalse(report["required"])
        self.assertEqual(report["status"], "not_applicable")


if __name__ == "__main__":
    unittest.main()
