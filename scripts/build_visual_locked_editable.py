from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def main() -> int:
    run = Path(sys.argv[1])
    output = Path(sys.argv[2])
    brief = json.loads((run / "slide_brief.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for index, item in enumerate(brief["slides"], 1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(run / "masters" / f"{item['slide_id']}.png"), 0, 0, width=prs.slide_width, height=prs.slide_height)
        exact = [item["title"], *item.get("content", [])]
        for text_index, text in enumerate(exact, 1):
            # Keep the visual layer pixel-identical while preserving reviewed
            # semantic text as editable PowerPoint objects in the Selection
            # Pane. These are intentionally off-canvas; the master remains the
            # release visual baseline until a human-reviewed placement pass.
            box = slide.shapes.add_textbox(Inches(13.4), Inches(0.05 + text_index * 0.12), Inches(2.0), Inches(0.1))
            box.name = f"semantic-text-{index:02d}-{text_index:02d}"
            tf = box.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            run_obj = p.add_run()
            run_obj.text = text
            run_obj.font.size = Pt(1)
            run_obj.font.color.rgb = RGBColor(255, 255, 255)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(json.dumps({"output": str(output.resolve()), "slides": len(prs.slides), "semantic_text_objects": sum(1 + len(s.get("content", [])) for s in brief["slides"])}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
