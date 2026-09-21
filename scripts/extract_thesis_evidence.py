from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: extract_thesis_evidence.py INPUT.docx OUTPUT.json", file=sys.stderr)
        return 2
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    doc = Document(src)
    paragraphs = []
    for idx, p in enumerate(doc.paragraphs):
        text = clean(p.text)
        if text:
            paragraphs.append({"index": idx, "style": p.style.name, "text": text})
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        tables.append({"index": ti, "rows": rows})
    headings = [p for p in paragraphs if p["style"].lower().startswith(("heading", "标题", "章", "节", "一级", "二级", "三级"))]
    # Include TOC-like paragraphs as structure even when Word styles are flattened.
    headings.extend([p for p in paragraphs if re.match(r"^(第[一二三四五六七八九十]+章|[1-9](?:\.[0-9]+){0,2})", p["text"])])
    seen = set()
    unique_headings = []
    for p in headings:
        if p["index"] not in seen:
            seen.add(p["index"])
            unique_headings.append(p)
    numbers = sorted(set(re.findall(r"(?<![A-Za-z])(?:\d+(?:\.\d+)?%?|\d+\.\d+|\d{3,})(?![A-Za-z])", "\n".join(p["text"] for p in paragraphs))))
    result = {
        "source": str(src),
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "paragraphs": paragraphs,
        "headings": unique_headings,
        "tables": tables,
        "candidate_numbers": numbers,
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("paragraph_count", "nonempty_paragraph_count", "table_count", "inline_shape_count", "section_count")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
