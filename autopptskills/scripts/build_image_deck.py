#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any


SLIDE_W = 13.333
SLIDE_H = 7.5
REQUIRED_OUTPUT_MODE = "imagegen_full_slide"
REQUIRED_GENERATION_MODE = "direct_final_slide_imagegen"


def rel_path(base: Path, path_value: str) -> Path:
    path = Path(path_value)
    if path.is_absolute():
        return path
    return (base / path).resolve()


def validate_spec(spec: dict[str, Any], spec_base: Path) -> list[Path]:
    deck = spec.get("deck", {})
    output_mode = deck.get("output_mode")
    generation_mode = deck.get("generation_mode")
    if output_mode != REQUIRED_OUTPUT_MODE:
        raise ValueError(
            "ImageGen assembly is imagegen-only. deck.output_mode must be "
            f"{REQUIRED_OUTPUT_MODE!r}; got {output_mode!r}."
        )
    if generation_mode != REQUIRED_GENERATION_MODE:
        raise ValueError(
            "ImageGen assembly is imagegen-only. deck.generation_mode must be "
            f"{REQUIRED_GENERATION_MODE!r}; got {generation_mode!r}."
        )

    slides = spec.get("slides", [])
    if not slides:
        raise ValueError("deck-spec.json must include at least one slide.")

    missing: list[Path] = []
    for item in slides:
        slide_id = item.get("slide_id", "<unknown>")
        if item.get("layout") != "full_slide_image":
            raise ValueError(
                "ImageGen assembly PPT assembly accepts only complete imagegen slides. "
                f"Slide {slide_id} must use layout='full_slide_image'."
            )
        slide_image = item.get("slide_image")
        if not slide_image:
            raise ValueError(f"Slide {slide_id} is missing slide_image.")
        image_path = rel_path(spec_base, slide_image)
        if not image_path.exists():
            missing.append(image_path)

    if missing:
        missing_list = "\n".join(f"- {path}" for path in missing)
        raise FileNotFoundError(
            "Missing generated full-slide image(s). ImageGen assembly has no alternate PPT route, "
            "no local text overlay path, and no diagnostic shell build.\n"
            f"{missing_list}\n"
            "Invoke Codex built-in image_gen for every pending slide, ingest the handoff, "
            "then run scripts/register_imagegen_outputs.py and this builder again."
        )

    return [rel_path(spec_base, item["slide_image"]) for item in slides]


def build_ppt(spec_path: Path, output_path: Path) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.util import Inches

    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    spec_base = spec_path.parent
    slide_images = validate_spec(spec, spec_base)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank_layout = prs.slide_layouts[6]

    for image_path in slide_images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(image_path), Inches(0), Inches(0), width=Inches(SLIDE_W), height=Inches(SLIDE_H))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return {
        "output": str(output_path),
        "slide_count": len(slide_images),
        "output_mode": REQUIRED_OUTPUT_MODE,
        "generation_mode": REQUIRED_GENERATION_MODE,
        "assembly_policy": "imagegen_full_slide_only",
        "slide_images": [str(path) for path in slide_images],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Assemble an ImageGen image-only PPTX from generated full-slide images."
    )
    parser.add_argument("deck_spec", help="Path to project/imagegen_ppt_generation/deck-spec.json")
    parser.add_argument("--output", required=True, help="Output PPTX path.")
    parser.add_argument("--report", default="", help="Optional JSON build report path.")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    result = build_ppt(
        Path(args.deck_spec).expanduser().resolve(),
        Path(args.output).expanduser().resolve(),
    )
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
