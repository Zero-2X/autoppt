"""Compose reviewed image measurements; preserve text and paths as native objects.

Input is a measured deck JSON. Raster crops require explicit bboxes and reasons.
Trace only isolated icon regions. No semantic full-slide image is embedded.
"""
from __future__ import annotations
import argparse, hashlib, json, subprocess
from pathlib import Path
import cv2
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('plan'); ap.add_argument('--out-dir', required=True)
    ap.add_argument('--node', default='node')
    args=ap.parse_args()
    plan=json.loads(Path(args.plan).read_text(encoding='utf-8'))
    out=Path(args.out_dir).resolve(); out.mkdir(parents=True, exist_ok=True)
    assets=out/'assets'; assets.mkdir(exist_ok=True)
    deck={k:v for k,v in plan.items() if k!='slides'}
    deck['slides']=[]; traces=[]; provenance=[]
    for i,s in enumerate(plan['slides']):
        src=Path(s['source']); im=Image.open(src).convert('RGB')
        if im.size!=(plan['ref_width'],plan['ref_height']):
            raise ValueError(f'{src}: source dimensions differ from measurements')
        slide={k:v for k,v in s.items() if k not in ('assets','source')}
        slide['icons']=[]
        for j,a in enumerate(s.get('assets',[])):
            x,y,w,h=map(int,a['source_bbox'])
            if not (x>=0 and y>=0 and w>0 and h>0 and x+w<=im.width and y+h<=im.height):
                raise ValueError(f'invalid source crop: {a}')
            crop=im.crop((x,y,x+w,y+h))
            p=assets/f'S{i+1:02d}-{j+1:02d}.png'; crop.save(p)
            record={**a,'source':str(src),'source_sha256':sha(src),'asset':str(p),'asset_sha256':sha(p)}
            provenance.append(record)
            if a.get('route')=='native-trace':
                rgb=np.array(crop)
                if a.get('ink')=='red':
                    mask=((rgb[:,:,0].astype(float)>rgb[:,:,1]*1.25)&(rgb[:,:,1]<150)).astype('uint8')*255
                elif a.get('ink')=='white':
                    mask=(rgb.min(axis=2)>220).astype('uint8')*255
                else:
                    mask=(rgb.max(axis=2)<135).astype('uint8')*255
                contours,hierarchy=cv2.findContours(mask,cv2.RETR_CCOMP,cv2.CHAIN_APPROX_SIMPLE)
                paths=[]
                if hierarchy is not None:
                    for k,c in enumerate(contours):
                        if cv2.contourArea(c)<2: continue
                        poly=cv2.approxPolyDP(c,0.35,True).reshape(-1,2)
                        paths.append({'points':poly.tolist(),'hole':int(hierarchy[0,k,3])>=0})
                if not paths or sum(len(p['points']) for p in paths)>6000:
                    raise ValueError(f'Unusable bounded trace: {a["name"]}')
                traces.append({'slide':i,**a,'paths':paths})
            else:
                if not a.get('reason'): raise ValueError('Raster exception needs a reason')
                slide['icons'].append({**a,'file':str(p),'x':x,'y':y,'w':w,'h':h,'layout_bbox':[x,y,w,h],'editability_level':'movable-image','role':'bounded-scientific-image','name':'scientific-image::'+a['name']})
        deck['slides'].append(slide)
    deckpath=out/'deck.json'; deckpath.write_text(json.dumps(deck,ensure_ascii=False,indent=2),encoding='utf-8')
    pptx=out/'defense_editable.pptx'
    subprocess.run([args.node,str(Path(__file__).with_name('compose_editable_pptx.mjs')),str(deckpath),str(pptx),'--report',str(out/'compose-report.json')],check=True)
    # Convert bounded traced contours to actual OOXML editable freeform paths.
    # python-pptx is used only for this native-freeform postprocess.
    prs=Presentation(pptx); sx=prs.slide_width/plan['ref_width']; sy=prs.slide_height/plan['ref_height']
    for t in traces:
        x,y,w,h=t['source_bbox']; slide=prs.slides[t['slide']]
        outer=next(p for p in t['paths'] if not p['hole'])
        p0=outer['points'][0]
        builder=slide.shapes.build_freeform(p0[0],p0[1],scale=(sx,sy))
        builder.add_line_segments(outer['points'][1:],close=True)
        shape=builder.convert_to_shape(origin_x=int(x*sx),origin_y=int(y*sy))
        shape.name='native-vector::'+t['name']
        shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(t.get('color','20242B'))
        shape.line.fill.background()
        geom=shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom')
        pathlist=geom.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pathLst')
        for ch in list(pathlist): pathlist.remove(ch)
        path=OxmlElement('a:path'); path.set('w',str(w)); path.set('h',str(h)); path.set('stroke','0')
        for poly in t['paths']:
            pts=poly['points']
            for k,(px,py) in enumerate(pts):
                command=OxmlElement('a:moveTo' if k==0 else 'a:lnTo')
                point=OxmlElement('a:pt'); point.set('x',str(px)); point.set('y',str(py)); command.append(point); path.append(command)
            path.append(OxmlElement('a:close'))
        pathlist.append(path)
        # Explicit extents use the measured crop, including contour margins.
        shape.left=int(x*sx); shape.top=int(y*sy); shape.width=int(w*sx); shape.height=int(h*sy)
    prs.save(pptx)
    (out/'asset-provenance.json').write_text(json.dumps({'assets':provenance,'native_trace_objects':len(traces),'trace_vertices':sum(len(p['points']) for t in traces for p in t['paths'])},ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'native-traces.json').write_text(json.dumps(traces,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'pptx':str(pptx),'slides':len(prs.slides),'native_traces':len(traces)}))

if __name__=='__main__': main()
