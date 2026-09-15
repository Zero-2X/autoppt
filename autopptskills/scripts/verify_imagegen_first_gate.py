#!/usr/bin/env python3
"""Verify that an image-only PPTX was assembled from real full-slide ImageGen outputs."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import struct
import sys
from pathlib import Path
from typing import Any, Iterable


REQUIRED_DECK_MODE = "imagegen_full_slide"
REQUIRED_GENERATION_MODE = "direct_final_slide_imagegen"
REQUIRED_POLICY = {
    "output_mode": REQUIRED_GENERATION_MODE,
    "imagegen_required": True,
    "builtin_imagegen_only": True,
    "external_api_key_allowed": False,
    "external_cli_allowed": False,
    "local_command_allowed": False,
    "mock_formal_output_allowed": False,
    "per_slide_imagegen_required": True,
    "complete_final_page_required": True,
    "no_native_ppt_overlay": True,
    "no_script_generated_slide_content": True,
}
OFFICIAL_SOURCE_PROFILE_IDS = {
    "graduate_defense_navy",
    "institutional_purple_light",
    "engineering_institutional_blue",
    "science_dark_contrast",
    "nsfc_review_light",
}
OFFICIAL_SOURCE_REQUIRED_POLICY = {
    "official_source_inspired_only": True,
    "forbid_official_template_claim": True,
    "forbid_institutional_identity_elements": True,
}
BUILTIN_BACKENDS = {"builtin", "builtin_image_gen"}
ACCEPTED_PIPELINE_STATUSES = {"generated", "skipped_existing_verified"}
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
IDENTITY_LIKE_EXACT_TEXT_RE = re.compile(
    r"(?:[\u4e00-\u9fffA-Za-z0-9·\- ]{1,48}(?:大学|学院|研究院|实验室)|"
    r"国家自然科学基金|\bNSFC\b|\bUniversity\b|\bCollege\b|\bSchool of\b|\bInstitute of\b)",
    re.IGNORECASE,
)


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def resolve_path(base: Path, value: str | Path) -> Path:
    path = Path(value).expanduser()
    return path.resolve() if path.is_absolute() else (base / path).resolve()


def validate_official_source_contract(
    deck: dict[str, Any],
    prompt_manifest: dict[str, Any],
    spec_slides: list[dict[str, Any]],
    prompt_slides: dict[str, dict[str, Any]],
) -> tuple[dict[str, Any], list[str]]:
    top_reference = prompt_manifest.get("style_reference")
    deck_reference = deck.get("style_reference")
    profile_ids = {
        str(deck.get("style_profile", "")),
        str(prompt_manifest.get("style_profile", "")),
        *(str(slide.get("style_profile", "")) for slide in prompt_slides.values()),
    }
    references = [
        top_reference,
        deck_reference,
        *(slide.get("style_reference") for slide in prompt_slides.values()),
    ]
    required = bool(
        OFFICIAL_SOURCE_PROFILE_IDS.intersection(profile_ids)
        or any(isinstance(ref, dict) and ref.get("official_source_inspired") is True for ref in references)
    )
    if not required:
        return {
            "required": False,
            "status": "not_applicable",
            "profile_id": str(prompt_manifest.get("style_profile") or deck.get("style_profile") or ""),
        }, []

    errors: list[str] = []
    profile_id = str(prompt_manifest.get("style_profile") or deck.get("style_profile") or "")
    if profile_id not in OFFICIAL_SOURCE_PROFILE_IDS:
        errors.append(f"official-source style_profile is missing or invalid: {profile_id!r}")
    if deck.get("style_profile") != profile_id:
        errors.append("deck and prompt manifest official-source style_profile values differ")

    if not isinstance(top_reference, dict):
        errors.append("official-source prompt manifest requires top-level style_reference")
        top_reference = {}
    if not isinstance(deck_reference, dict):
        errors.append("official-source deck spec requires deck.style_reference")
    elif deck_reference != top_reference:
        errors.append("deck.style_reference differs from prompt manifest style_reference")

    if top_reference.get("reference_mode") != "official-source-inspired-generic-archetype":
        errors.append("official-source style_reference.reference_mode is invalid")
    if top_reference.get("official_source_inspired") is not True:
        errors.append("official-source style_reference.official_source_inspired must be true")
    if top_reference.get("official_template") is not False:
        errors.append("official-source style_reference.official_template must be false")
    sources = top_reference.get("sources")
    if not isinstance(sources, list) or not sources:
        errors.append("official-source style_reference.sources must be non-empty")
    elif any(
        not isinstance(source, dict) or source.get("asset_reuse_allowed") is not False
        for source in sources
    ):
        errors.append("official-source style_reference sources must all set asset_reuse_allowed=false")
    if not isinstance(top_reference.get("allowed_extraction"), list) or not top_reference["allowed_extraction"]:
        errors.append("official-source style_reference.allowed_extraction must be non-empty")
    if not isinstance(top_reference.get("forbidden_identity_elements"), list) or not top_reference["forbidden_identity_elements"]:
        errors.append("official-source style_reference.forbidden_identity_elements must be non-empty")
    if not str(top_reference.get("identity_guard", "")).strip():
        errors.append("official-source style_reference.identity_guard must be non-empty")
    if top_reference.get("identity_text_policy") != "source_verified_exact_text_only_not_identity_asset":
        errors.append("official-source style_reference.identity_text_policy is invalid")
    verified_identity_text = top_reference.get("verified_identity_text", [])
    allowed_identity_alias = top_reference.get("allowed_identity_text", [])
    if not isinstance(verified_identity_text, list) or not all(
        isinstance(item, str) and item.strip() for item in verified_identity_text
    ):
        errors.append("official-source style_reference.verified_identity_text must be a list of non-empty strings")
        verified_identity_text = []
    if not isinstance(allowed_identity_alias, list) or not all(
        isinstance(item, str) and item.strip() for item in allowed_identity_alias
    ):
        errors.append("official-source style_reference.allowed_identity_text must be a list of non-empty strings")
        allowed_identity_alias = []
    if verified_identity_text != allowed_identity_alias:
        errors.append("verified_identity_text and allowed_identity_text must match")
    allowed_identity_text = set(verified_identity_text)
    exact_text = {
        str(item).strip()
        for slide in prompt_slides.values()
        for item in (slide.get("exact_text") or [])
        if str(item).strip()
    }
    missing_identity_text = sorted(allowed_identity_text - exact_text)
    if missing_identity_text:
        errors.append(
            "verified identity text does not exactly match any slide exact_text item: "
            + "; ".join(missing_identity_text)
        )
    unapproved_identity_text = sorted(
        item
        for item in exact_text
        if IDENTITY_LIKE_EXACT_TEXT_RE.search(item) and item not in allowed_identity_text
    )
    if unapproved_identity_text:
        errors.append(
            "identity-like slide exact_text is not explicitly source-verified: "
            + "; ".join(unapproved_identity_text)
        )

    policy = prompt_manifest.get("prompt_policy", {})
    for key, expected in OFFICIAL_SOURCE_REQUIRED_POLICY.items():
        if policy.get(key) != expected:
            errors.append(f"prompt_policy.{key} must be {expected!r} for an official-source profile")

    for spec_slide in spec_slides:
        slide_id = str(spec_slide.get("slide_id", ""))
        prompt_slide = prompt_slides.get(slide_id)
        if not prompt_slide:
            continue
        if prompt_slide.get("style_profile") != profile_id:
            errors.append(f"{slide_id}: style_profile differs from the official-source deck profile")
        if prompt_slide.get("style_reference") != top_reference:
            errors.append(f"{slide_id}: style_reference differs from the top-level official-source contract")

    return {
        "required": True,
        "status": "pass" if not errors else "fail",
        "profile_id": profile_id,
        "verified_identity_text": list(verified_identity_text),
        "required_policy": dict(OFFICIAL_SOURCE_REQUIRED_POLICY),
    }, errors


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def png_dimensions(path: Path) -> tuple[int, int] | None:
    with path.open("rb") as handle:
        header = handle.read(24)
    if len(header) < 24 or not header.startswith(PNG_SIGNATURE):
        return None
    return struct.unpack(">II", header[16:24])


def iter_dicts(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from iter_dicts(child)
    elif isinstance(value, list):
        for child in value:
            yield from iter_dicts(child)


def first_string(record: dict[str, Any], keys: Iterable[str]) -> str:
    for key in keys:
        value = record.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def classify_provenance(record: dict[str, Any], root: dict[str, Any]) -> tuple[bool, str]:
    imagegen_id = first_string(record, ("id", "imagegen_id", "generation_id"))
    backend = first_string(record, ("backend", "imagegen_backend")).lower()
    status = first_string(record, ("status", "provenance_status")).lower()
    mode = first_string(record, ("generation_mode",)) or first_string(root, ("generation_mode",))
    provenance_kind = first_string(record, ("provenance_kind",)).lower()
    builtin_ok = (
        mode == REQUIRED_GENERATION_MODE
        and status == "generated"
        and (
            backend in BUILTIN_BACKENDS
            or provenance_kind in {"builtin-imagegen", "builtin_image_gen", "image_generation_call"}
            or imagegen_id.startswith("ig_")
        )
    )
    return builtin_ok, "strong-ig-id" if builtin_ok and imagegen_id.startswith("ig_") else ("builtin-imagegen" if builtin_ok else "unverified")


def provenance_records(paths: list[Path]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    json_paths: list[Path] = []
    for path in paths:
        if path.is_dir():
            json_paths.extend(sorted(path.rglob("*.json")))
        elif path.is_file():
            json_paths.append(path)

    seen: set[tuple[str, str, str, str]] = set()
    for path in json_paths:
        try:
            data = load_json(path)
        except (OSError, json.JSONDecodeError):
            continue
        root = data if isinstance(data, dict) else {}
        for item in iter_dicts(data):
            digest = first_string(item, ("sha256", "image_sha256", "output_sha256")).lower()
            if not SHA256_RE.fullmatch(digest):
                continue
            accepted, strength = classify_provenance(item, root)
            imagegen_id = first_string(item, ("id", "imagegen_id", "generation_id"))
            backend = first_string(item, ("backend", "imagegen_backend"))
            key = (str(path.resolve()), digest, imagegen_id, backend)
            if key in seen:
                continue
            seen.add(key)
            records.append(
                {
                    "sha256": digest,
                    "accepted": accepted,
                    "strength": strength,
                    "imagegen_id": imagegen_id,
                    "backend": backend,
                    "status": first_string(item, ("status", "provenance_status")),
                    "record_index": item.get("record_index"),
                    "source": str(path.resolve()),
                }
            )
    return records


def has_full_slide_prompt_contract(slide: dict[str, Any]) -> bool:
    expected = str(slide.get("expected_output", "")).lower()
    criteria = "\n".join(str(value) for value in slide.get("acceptance_criteria", []))
    prompt = "\n".join(
        str(slide.get(key, "")) for key in ("prompt_zh", "prompt_en", "prompt")
    ).lower()
    chinese_complete_page_contract = "生成一张完整、最终可用的中文powerpoint幻灯片图片" in prompt
    # Older onlyppt prompt packs already used an explicit Chinese 16:9
    # complete-slide contract but did not duplicate it into expected_output.
    # Preserve those manifests instead of forcing a prompt rewrite or a new
    # ImageGen call; the top-level prompt policy still carries the hard
    # complete_final_page_required invariant.
    legacy_chinese_complete_page_contract = "16:9 完整中文" in prompt
    expected_ok = (
        "complete final ppt page image" in expected
        or "complete final ppt page image" in criteria.lower()
        or chinese_complete_page_contract
        or legacy_chinese_complete_page_contract
    )
    prompt_ok = any(
        phrase in prompt
        for phrase in (
            "complete final 16:9 powerpoint slide image, one page only",
            "whole output must be one finished slide page",
            "complete final ppt page image",
        )
    ) or chinese_complete_page_contract or legacy_chinese_complete_page_contract
    return expected_ok and prompt_ok


def inspect_image_pptx(path: Path) -> tuple[dict[str, Any], list[str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    errors: list[str] = []
    prs = Presentation(str(path))
    totals = {
        "slides": len(prs.slides),
        "pictures": 0,
        "full_slide_pictures": 0,
        "text_boxes": 0,
        "native_shapes": 0,
    }
    slide_details = []
    tolerance = max(prs.slide_width, prs.slide_height) * 0.005

    for index, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        text_boxes = [shape for shape in shapes if getattr(shape, "has_text_frame", False)]
        non_pictures = [shape for shape in shapes if shape.shape_type != MSO_SHAPE_TYPE.PICTURE]
        full_slide = [
            shape
            for shape in pictures
            if abs(shape.left) <= tolerance
            and abs(shape.top) <= tolerance
            and abs(shape.width - prs.slide_width) <= tolerance
            and abs(shape.height - prs.slide_height) <= tolerance
        ]
        totals["pictures"] += len(pictures)
        totals["full_slide_pictures"] += len(full_slide)
        totals["text_boxes"] += len(text_boxes)
        totals["native_shapes"] += len(non_pictures)
        slide_details.append(
            {
                "slide": index,
                "objects": len(shapes),
                "pictures": len(pictures),
                "full_slide_pictures": len(full_slide),
                "text_boxes": len(text_boxes),
                "native_shapes": len(non_pictures),
            }
        )
        if len(shapes) != 1 or len(pictures) != 1 or len(full_slide) != 1:
            errors.append(
                f"PPT slide {index} must contain exactly one full-slide picture and no other objects; "
                f"found objects={len(shapes)}, pictures={len(pictures)}, full_slide={len(full_slide)}."
            )

    return {"pptx": str(path.resolve()), **totals, "slide_details": slide_details}, errors


def run_gate(args: argparse.Namespace) -> dict[str, Any]:
    run_dir = Path(args.run_dir).expanduser().resolve()
    deck_spec_path = resolve_path(run_dir, args.deck_spec or "deck-spec.json")
    prompt_manifest_path = resolve_path(run_dir, args.prompt_manifest or "image-prompts.json")
    pptx_path = Path(args.pptx).expanduser().resolve()
    provenance_paths = (
        [Path(value).expanduser().resolve() for value in args.provenance]
        if args.provenance
        else [run_dir / "assets" / "generated", run_dir / "references" / "asset-manifest.json"]
    )

    errors: list[str] = []
    warnings: list[str] = []
    for required in (deck_spec_path, prompt_manifest_path, pptx_path):
        if not required.exists():
            errors.append(f"Missing required gate input: {required}")
    if errors:
        return {"gate": "imagegen-first-full-slide", "verdict": "fail", "errors": errors, "warnings": warnings}

    deck_spec = load_json(deck_spec_path)
    prompt_manifest = load_json(prompt_manifest_path)
    deck = deck_spec.get("deck", {})
    if deck.get("output_mode") != REQUIRED_DECK_MODE:
        errors.append(f"deck.output_mode must be {REQUIRED_DECK_MODE!r}.")
    if deck.get("generation_mode") != REQUIRED_GENERATION_MODE:
        errors.append(f"deck.generation_mode must be {REQUIRED_GENERATION_MODE!r}.")

    policy = prompt_manifest.get("prompt_policy", {})
    for key, expected in REQUIRED_POLICY.items():
        if policy.get(key) != expected:
            errors.append(f"prompt_policy.{key} must be {expected!r}; got {policy.get(key)!r}.")

    spec_slides = deck_spec.get("slides", [])
    prompt_slides = {slide.get("slide_id"): slide for slide in prompt_manifest.get("slides", [])}
    official_source_contract, official_source_errors = validate_official_source_contract(
        deck,
        prompt_manifest,
        spec_slides,
        prompt_slides,
    )
    errors.extend(official_source_errors)
    records = provenance_records(provenance_paths)
    by_hash: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        by_hash.setdefault(record["sha256"], []).append(record)

    slide_reports = []
    verified_count = 0
    strong_count = 0
    contract_count = 0
    dimensions_seen: set[tuple[int, int]] = set()
    for spec_slide in spec_slides:
        slide_id = str(spec_slide.get("slide_id", ""))
        if spec_slide.get("layout") != "full_slide_image":
            errors.append(f"{slide_id or '<unknown>'}: layout must be 'full_slide_image'.")
        image_value = spec_slide.get("slide_image")
        image_path = resolve_path(deck_spec_path.parent, image_value) if image_value else Path()
        prompt_slide = prompt_slides.get(slide_id)
        contract_ok = bool(prompt_slide and has_full_slide_prompt_contract(prompt_slide))
        if contract_ok:
            contract_count += 1
        else:
            errors.append(f"{slide_id}: prompt does not contract for one complete final PPT page image.")

        if not image_value or not image_path.exists():
            errors.append(f"{slide_id}: missing final slide image {image_path}.")
            slide_reports.append({"slide_id": slide_id, "verified": False, "png": str(image_path)})
            continue

        digest = sha256_file(image_path)
        dimensions = png_dimensions(image_path)
        if dimensions:
            dimensions_seen.add(dimensions)
        matches = [record for record in by_hash.get(digest, []) if record["accepted"]]
        strong_matches = [record for record in matches if record["strength"] == "strong-ig-id"]
        verified = bool(strong_matches if args.require_strong_ig_id else matches)
        if verified:
            verified_count += 1
            if strong_matches:
                strong_count += 1
        else:
            requirement = "a matching ig_ ImageGen record" if args.require_strong_ig_id else "matching ImageGen provenance"
            errors.append(f"{slide_id}: SHA-256 {digest} has no {requirement}.")

        chosen = (strong_matches or matches or by_hash.get(digest, []) or [{}])[0]
        slide_reports.append(
            {
                "slide_id": slide_id,
                "verified": verified,
                "png": str(image_path),
                "sha256": digest,
                "dimensions": list(dimensions) if dimensions else None,
                "provenance_json": chosen.get("source", ""),
                "imagegen_id": chosen.get("imagegen_id", ""),
                "backend": chosen.get("backend", ""),
                "provenance_strength": chosen.get("strength", "unverified"),
                "manifest_full_slide_contract": contract_ok,
            }
        )

    if set(prompt_slides) != {str(slide.get("slide_id", "")) for slide in spec_slides}:
        errors.append("Slide IDs differ between deck-spec.json and image-prompts.json.")

    assembly, pptx_errors = inspect_image_pptx(pptx_path)
    errors.extend(pptx_errors)
    if assembly["slides"] != len(spec_slides):
        errors.append(
            f"PPTX slide count {assembly['slides']} does not match deck spec count {len(spec_slides)}."
        )

    if len(dimensions_seen) > 1:
        warnings.append(
            "ImageGen returned multiple native PNG dimensions. This is acceptable only because the image-only "
            "assembler places each verified complete page full-slide; inspect stretching/cropping visually."
        )

    verdict = "pass" if not errors else "fail"
    return {
        "gate": "imagegen-first-full-slide",
        "run_dir": str(run_dir),
        "slides_expected": len(spec_slides),
        "slides_verified": verified_count,
        "strong_ig_id_slides": strong_count,
        "manifest_full_slide_contracts": contract_count,
        "current_png_dimensions": [list(item) for item in sorted(dimensions_seen)],
        "assembly": {**assembly, "build_policy": "imagegen_full_slide_only"},
        "proof_logic": [
            "each current slide PNG SHA-256 matches accepted ImageGen provenance",
            "each slide manifest explicitly requests one complete final PPT page image",
            "the image-only PPTX contains exactly one full-slide picture per slide and zero other objects",
        ],
        "script_role": (
            "verify provenance and inspect assembly structure only; the Codex host invokes built-in image_gen, "
            "repository scripts only ingest/assemble and never call a provider"
        ),
        "require_strong_ig_id": bool(args.require_strong_ig_id),
        "official_source_inspired": bool(official_source_contract.get("required")),
        "official_source_contract": official_source_contract,
        "verdict": verdict,
        "errors": errors,
        "warnings": warnings,
        "slides": slide_reports,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_dir", help="Run/stage directory containing deck-spec.json and image-prompts.json.")
    parser.add_argument("--pptx", required=True, help="Image-only PPTX to inspect.")
    parser.add_argument("--deck-spec", default="", help="Override deck-spec.json path, relative to run_dir.")
    parser.add_argument("--prompt-manifest", default="", help="Override image-prompts.json path, relative to run_dir.")
    parser.add_argument(
        "--provenance",
        action="append",
        default=[],
        help="Provenance JSON file or directory; repeat as needed. Defaults to assets/generated and references/asset-manifest.json.",
    )
    parser.add_argument(
        "--require-strong-ig-id",
        action="store_true",
        help="Require every current slide hash to match a built-in ImageGen record whose id starts with ig_.",
    )
    parser.add_argument("--report", default="", help="Optional JSON report output path.")
    return parser


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = build_parser().parse_args()
    report = run_gate(args)
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report.get("verdict") == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
