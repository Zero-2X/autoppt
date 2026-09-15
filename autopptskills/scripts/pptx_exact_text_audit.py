#!/usr/bin/env python3
"""Audit source exact-text tokens against native text and disclosed text fallbacks.

The normal gold-release contract still prefers native PowerPoint text.  Some
high-fidelity reconstructions intentionally keep a bounded, movable image for
art-directed typography when native text visibly drifts from the source.  When
``--layout`` is supplied, those explicitly declared ``art-text-fallback``
objects are audited as a second evidence channel, never silently counted as
native editability.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from pptx import Presentation


SEPARATOR_RE = re.compile(r"[\u00b7\u2022\u2192\uff5c|\n]+")
TOKEN_RE = re.compile(
    r"[A-Za-z]+[0-9]+(?:\.[0-9]+)?"
    r"|[A-Za-z]+(?:-[A-Za-z]+)*"
    r"|[0-9]+(?:\.[0-9]+)?(?:\s*[~\-]\s*[0-9]+(?:\.[0-9]+)?)?%?"
    r"|[\u3400-\u4dbf\u4e00-\u9fff]+"
)


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKC", str(value)).lower().replace("\u2103", "c")
    return re.sub(r"[^0-9a-z%+\-.\u3400-\u4dbf\u4e00-\u9fff]+", "", value)


def meaningful_tokens(value: str) -> list[str]:
    tokens: list[str] = []
    for segment in SEPARATOR_RE.split(str(value)):
        segment = segment.strip(" \t\r\n,.;\uff0c\u3002\uff1b\u3001")
        if not segment:
            continue
        parts = re.split(r"[\uff1a:]", segment, maxsplit=1)
        for part in parts:
            for match in TOKEN_RE.finditer(part):
                token = normalize(match.group(0))
                if len(token) >= 2 and token not in tokens:
                    tokens.append(token)
    return tokens


def slide_native_text(slide: Any) -> str:
    return "\n".join(
        str(getattr(shape, "text", "") or "")
        for shape in slide.shapes
        if str(getattr(shape, "text", "") or "").strip()
    )


def load_layout_fallbacks(layout: Path | None) -> dict[str, list[str]]:
    """Return disclosed movable art-text strings grouped by slide id."""
    if layout is None:
        return {}
    payload = json.loads(layout.read_text(encoding="utf-8"))
    fallbacks: dict[str, list[str]] = {}
    for item in payload.get("slides", []):
        slide_id = str(item.get("slide_id", ""))
        if not slide_id:
            continue
        values: list[str] = []
        for icon in item.get("icons", []) or []:
            if str(icon.get("role", "")) != "art-text-fallback":
                continue
            text = str(icon.get("text", "") or "").strip()
            if text:
                values.append(text)
        fallbacks[slide_id] = values
    return fallbacks


def audit(pptx: Path, manifest: Path, layout: Path | None = None) -> dict[str, Any]:
    prs = Presentation(str(pptx))
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    raw_slides = payload if isinstance(payload, list) else payload.get("slides", [])
    manifest_slides = {str(item["slide_id"]): item for item in raw_slides}
    layout_fallbacks = load_layout_fallbacks(layout)
    missing: list[dict[str, Any]] = []
    reference_reports = []
    total_tokens = 0
    matched_tokens = 0
    native_matched_tokens = 0
    fallback_matched_tokens = 0
    fallback_references = 0

    for slide_number, slide in enumerate(prs.slides, start=1):
        slide_id = f"S{slide_number:02d}"
        native_text = slide_native_text(slide)
        normalized_native = normalize(native_text)
        fallback_text = "\n".join(layout_fallbacks.get(slide_id, []))
        normalized_fallback = normalize(fallback_text)
        for reference_index, exact_text in enumerate(
            manifest_slides.get(slide_id, {}).get("exact_text", [])
        ):
            tokens = meaningful_tokens(str(exact_text))
            native_tokens = [token for token in tokens if token in normalized_native]
            fallback_tokens = [
                token for token in tokens
                if token not in normalized_native and token in normalized_fallback
            ]
            missing_tokens = [
                token for token in tokens
                if token not in normalized_native and token not in normalized_fallback
            ]
            total_tokens += len(tokens)
            native_matched_tokens += len(native_tokens)
            fallback_matched_tokens += len(fallback_tokens)
            matched_tokens += len(native_tokens) + len(fallback_tokens)
            if fallback_tokens:
                fallback_references += 1
            report = {
                "slide_id": slide_id,
                "reference_index": reference_index,
                "exact_text": str(exact_text),
                "tokens": tokens,
                "native_tokens": native_tokens,
                "fallback_tokens": fallback_tokens,
                "missing_tokens": missing_tokens,
            }
            reference_reports.append(report)
            if missing_tokens:
                missing.append(report)

    coverage = matched_tokens / max(1, total_tokens)
    native_coverage = native_matched_tokens / max(1, total_tokens)
    fallback_coverage = fallback_matched_tokens / max(1, total_tokens)
    return {
        "pptx": str(pptx.resolve()),
        "manifest": str(manifest.resolve()),
        "layout": str(layout.resolve()) if layout else None,
        "slides": len(prs.slides),
        "references": len(reference_reports),
        "tokens": total_tokens,
        "matched_tokens": matched_tokens,
        "native_matched_tokens": native_matched_tokens,
        "fallback_matched_tokens": fallback_matched_tokens,
        "missing_tokens": total_tokens - matched_tokens,
        "token_coverage": round(coverage, 4),
        "native_token_coverage": round(native_coverage, 4),
        "fallback_token_coverage": round(fallback_coverage, 4),
        "fallback_references": fallback_references,
        "fallback_enabled": bool(layout),
        "missing_references": len(missing),
        "missing": missing,
        "verdict": "pass" if not missing else "fail",
        "native_verdict": "pass" if native_matched_tokens == total_tokens else "fail",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx")
    parser.add_argument("manifest")
    parser.add_argument(
        "--layout",
        help=(
            "Optional deck-separated-editable.json. Explicit art-text-fallback "
            "strings are audited separately from native PowerPoint text."
        ),
    )
    parser.add_argument("--json-out")
    args = parser.parse_args()

    report = audit(
        Path(args.pptx),
        Path(args.manifest),
        Path(args.layout) if args.layout else None,
    )
    text = json.dumps(report, ensure_ascii=False, indent=2)
    print(text)
    if args.json_out:
        out = Path(args.json_out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text + "\n", encoding="utf-8")
    return 0 if report["verdict"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
